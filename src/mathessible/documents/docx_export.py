"""A Markdown document as a Word document (DOCX), every equation a native Word equation.

The Markdown goes through markdown-it, each equation through the
converter (MathML) and `omml_math`, written into the paragraph as Office
Math: inline in its line, a display equation as a paragraph of its own.
Headings use Word's Heading styles (navigation pane, table of contents,
screen readers), lists its List styles, footnotes are Word's own, the
document's header and footer (its front matter) are the page's, the
document's properties and language are set, and the page, font, page
numbers and a table of contents come from `DocxOptions`.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import unquote
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph

from mathessible.documents.charts import ChartError, alternative_text, draw, is_chart, parse_chart
from mathessible.documents.flowcharts import (
    FlowchartError,
    is_flowchart,
    parse_flowchart,
)
from mathessible.documents.flowcharts import (
    alternative_text as flowchart_alternative_text,
)
from mathessible.documents.flowcharts import (
    draw as draw_flowchart,
)
from mathessible.documents.front_matter import HeaderFooter, header_and_footer
from mathessible.documents.front_matter import body as without_front_matter
from mathessible.documents.html_export import ExportPage, for_page, markdown_parser
from mathessible.documents.markdown_math import find_math_zones, with_marks
from mathessible.documents.omml_math import M, OmmlEquation

# Paper sizes in millimetres, portrait.
PAPERS = {"a4": (210, 297), "a5": (148, 210), "a3": (297, 420), "us-letter": (215.9, 279.4), "us-legal": (215.9, 355.6)}
_MONOSPACE = "Courier New"
# The pictures Word shows, which an image beside the document may be (not SVG).
_PICTURES = (".png", ".jpg", ".jpeg", ".gif")


@dataclass(frozen=True, slots=True)
class DocxOptions:
    paper: str = "a4"
    landscape: bool = False
    margins_mm: tuple[float, float, float, float] = (25, 25, 25, 25)  # Top, right, bottom, left.
    font: str = "Calibri"
    font_size: float = 11
    line_spacing: float = 1.15  # A multiple of single spacing.
    justify: bool = False
    page_numbers: bool = True
    table_of_contents: bool = False
    title: str = ""
    author: str = ""
    subject: str = ""
    keywords: tuple[str, ...] = ()
    language: str = "en-US"  # Word's language tag, for proofing and screen readers.


# An equation's Office Math; None when it does not build.
OmmlOf = Callable[[str], OmmlEquation | None]


def _placeholder(index: int) -> str:
    return f"{index}"


_PLACEHOLDER = re.compile("(\\d+)")


def _language_tag(language: str) -> str:
    """Word wants a region: "ro" → "ro-RO", "en" → "en-US"."""
    if "-" in language:
        return language
    return {"en": "en-US", "pt": "pt-PT", "zh": "zh-CN"}.get(language, f"{language}-{language.upper()}")


class _Footnotes:
    """Word's footnotes part (`word/footnotes.xml`), made when the first note is written."""

    def __init__(self, document: Document):
        self.document = document
        self.part: XmlPart | None = None
        self.count = 0

    def add(self) -> tuple[int, Paragraph]:
        """A new footnote: its id and its first paragraph, which starts with the note's number."""
        if self.part is None:
            self._create()
        self.count += 1
        note = parse_xml(f'<w:footnote {nsdecls("w")} w:id="{self.count}"/>')
        self.part.element.append(note)
        paragraph = Paragraph(OxmlElement("w:p"), self)
        note.append(paragraph._p)
        paragraph._p.style = "footnotetext"  # The part has no styles of its own to look the name up in.
        _reference_run(paragraph)._r.append(OxmlElement("w:footnoteRef"))
        paragraph.add_run(" ")
        return self.count, paragraph

    def paragraph(self, note: int) -> Paragraph:
        """Another paragraph of the note."""
        paragraph = Paragraph(OxmlElement("w:p"), self)
        self.part.element[-1].append(paragraph._p)
        paragraph._p.style = "footnotetext"  # The part has no styles of its own to look the name up in.
        return paragraph

    def _create(self) -> None:
        separators = "".join(
            f'<w:footnote w:type="{kind}" w:id="{number}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
            f'w:lineRule="auto"/></w:pPr><w:r><w:{kind}/></w:r></w:p></w:footnote>'
            for number, kind in ((-1, "separator"), (0, "continuationSeparator"))
        )
        element = parse_xml(f'<w:footnotes {nsdecls("w", "r")} xmlns:m="{M}">{separators}</w:footnotes>')
        package = self.document.part.package
        self.part = XmlPart(PackURI("/word/footnotes.xml"), CONTENT_TYPE.WML_FOOTNOTES, element, package)
        self.document.part.relate_to(self.part, RELATIONSHIP_TYPE.FOOTNOTES)
        styles = self.document.styles
        text = styles.add_style("footnote text", WD_STYLE_TYPE.PARAGRAPH)
        text.base_style = styles["Normal"]
        text.font.size = Pt(10)
        text.paragraph_format.space_after = Pt(0)
        text.paragraph_format.line_spacing = 1
        reference = styles.add_style("footnote reference", WD_STYLE_TYPE.CHARACTER)
        reference.font.superscript = True


def _reference_run(paragraph: Paragraph) -> object:
    """A run in the footnote reference style (superscript), which `_Footnotes` adds with the first note."""
    run = paragraph.add_run()
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "footnotereference")
    run._r.get_or_add_rPr().insert(0, style)
    return run


class _Writer:
    def __init__(
        self, document: Document, markdown: str, omml_of: OmmlOf, footnotes: _Footnotes, folder: Path | None = None
    ):
        self.document = document
        self.folder = folder  # Where the document's images are.
        markdown = without_front_matter(markdown)
        self.zones = find_math_zones(markdown)
        self.sources = [zone.content(markdown) for zone in self.zones]
        self.omml_of = omml_of
        self.footnotes = footnotes
        self.tokens = markdown_parser().parse(with_marks(markdown, self.zones, _placeholder))
        self.notes: dict[int, list[object]] = {}  # Each note's inline tokens, one per paragraph.
        self.numbers: dict[int, int] = {}  # Word's number of each note written.
        current: int | None = None
        for token in self.tokens:
            if token.type == "footnote_open":
                current = token.meta["id"]
                self.notes[current] = []
            elif token.type == "footnote_close":
                current = None
            elif token.type == "inline" and current is not None:
                self.notes[current].append(token)

    def _footnote(self, paragraph: Paragraph, meta: dict) -> None:
        """A footnote reference; the note itself is written into Word's footnotes the first time."""
        note = meta["id"]
        run = _reference_run(paragraph)
        if note in self.numbers:
            run.add_text(str(self.numbers[note]))  # Referenced again: its number, as Word's NOTEREF shows it.
            return
        number, first = self.footnotes.add()
        self.numbers[note] = number
        reference = OxmlElement("w:footnoteReference")
        reference.set(qn("w:id"), str(number))
        run._r.append(reference)
        for index, token in enumerate(self.notes.get(note, [])):
            self.inline(first if index == 0 else self.footnotes.paragraph(number), token)

    # Inline content

    def _equation(self, paragraph: object, index: int) -> None:
        source = self.sources[index]
        equation = self.omml_of(source)
        if equation is None:
            run = paragraph.add_run(source)  # Does not build: its UnicodeMath, as code.
            run.font.name = _MONOSPACE
            return
        paragraph._p.append(parse_xml(equation.xml))

    def _text(self, paragraph: object, text: str, style: dict[str, bool], link: str | None) -> None:
        position = 0
        for match in _PLACEHOLDER.finditer(text):
            self._run(paragraph, text[position : match.start()], style, link)
            self._equation(paragraph, int(match.group(1)))
            position = match.end()
        self._run(paragraph, text[position:], style, link)

    def _run(self, paragraph: object, text: str, style: dict[str, bool], link: str | None) -> None:
        if not text:
            return
        if link:
            self._hyperlink(paragraph, text, link)
            return
        run = paragraph.add_run(text)
        run.bold = style.get("bold") or None
        run.italic = style.get("italic") or None
        if style.get("strike"):
            run.font.strike = True
        if style.get("code"):
            run.font.name = _MONOSPACE

    def _hyperlink(self, paragraph: object, text: str, url: str) -> None:
        relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        style = OxmlElement("w:rStyle")
        style.set(qn("w:val"), "Hyperlink")
        properties.append(style)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(underline)
        run.append(properties)
        text_element = OxmlElement("w:t")
        text_element.text = text
        text_element.set(qn("xml:space"), "preserve")
        run.append(text_element)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    @staticmethod
    def _describe(run: object, description: str) -> None:
        """A picture's alternative text, which Word reads out (its docPr's descr)."""
        for properties in run._r.xpath(".//wp:docPr"):
            properties.set("descr", description)

    def _image(self, paragraph: object, token: object, style: dict, link: str | None) -> None:
        """An image beside the document, its description its alternative text; else that description."""
        source = token.attrGet("src") or ""
        path = (self.folder / unquote(source)) if self.folder is not None and "://" not in source else None
        if path is None or path.suffix.lower() not in _PICTURES or not path.is_file():
            self._run(paragraph, token.content, style, link)
            return
        run = paragraph.add_run()
        try:
            run.add_picture(str(path), width=Mm(140))
        except (OSError, ValueError):
            self._run(paragraph, token.content, style, link)
            return
        self._describe(run, token.content)

    def _chart(self, text: str) -> None:
        """A chart drawn (charts.draw, as PNG: Word has no SVG of its own), its alternative text the chart in
        words, its title as Word's caption; its source as code when it cannot be drawn."""
        try:
            chart = parse_chart(text)
            picture = draw(chart, "png")
        except ChartError:
            for line in text.rstrip("\n").split("\n"):
                self.document.add_paragraph(style="No Spacing").add_run(line).font.name = _MONOSPACE
            return
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(io.BytesIO(picture), width=Mm(150))
        self._describe(run, alternative_text(chart))
        if chart.title:
            self.document.add_paragraph(chart.title, style="Caption")

    def _flowchart(self, text: str) -> None:
        """A Mermaid flowchart drawn locally as PNG, with alternative text and an optional caption."""
        try:
            flowchart = parse_flowchart(text)
            picture = draw_flowchart(flowchart, "png")
        except FlowchartError:
            for line in text.rstrip("\n").split("\n"):
                self.document.add_paragraph(style="No Spacing").add_run(line).font.name = _MONOSPACE
            return
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(io.BytesIO(picture), width=Mm(150))
        self._describe(run, flowchart_alternative_text(flowchart))
        if flowchart.title:
            self.document.add_paragraph(flowchart.title, style="Caption")

    def inline(self, paragraph: object, token: object) -> None:
        style = {"bold": False, "italic": False, "strike": False, "code": False}
        link: str | None = None
        for child in token.children or ():
            kind = child.type
            if kind == "text":
                self._text(paragraph, child.content, style, link)
            elif kind == "code_inline":
                self._run(paragraph, child.content, {**style, "code": True}, link)
            elif kind == "softbreak":
                self._run(paragraph, " ", style, link)
            elif kind == "hardbreak":
                paragraph.add_run().add_break()
            elif kind in ("strong_open", "strong_close"):
                style["bold"] = kind == "strong_open"
            elif kind in ("em_open", "em_close"):
                style["italic"] = kind == "em_open"
            elif kind in ("s_open", "s_close"):
                style["strike"] = kind == "s_open"
            elif kind == "link_open":
                link = child.attrGet("href") or None
            elif kind == "link_close":
                link = None
            elif kind == "image":
                self._image(paragraph, child, style, link)
            elif kind == "footnote_ref":
                self._footnote(paragraph, child.meta)

    # Blocks

    def _display(self, index: int) -> None:
        """A display equation: Word's equation paragraph (`m:oMathPara`), centred; its number after it."""
        source = self.sources[index]
        equation = self.omml_of(source)
        paragraph = self.document.add_paragraph()
        if equation is None:
            run = paragraph.add_run(source)
            run.font.name = _MONOSPACE
            return
        math = parse_xml(equation.xml)
        if equation.number:
            # As Word numbers an equation (`#` in its linear format): an array whose line ends in "#(1)".
            number = escape(equation.number.strip("()"))
            math = parse_xml(
                f'<m:oMath xmlns:m="{M}"><m:eqArr><m:e>{equation.body}<m:r><m:t>#</m:t></m:r>'
                f"<m:d><m:e><m:r><m:t>{number}</m:t></m:r></m:e></m:d></m:e></m:eqArr></m:oMath>"
            )
        block = OxmlElement("m:oMathPara")
        block.append(math)
        paragraph._p.append(block)

    def body(self) -> None:
        lists: list[str] = []  # "Bullet" or "Number", per open list.
        quote = 0
        heading = 0
        table: list[list[object]] | None = None
        cell_token: list[object] = []
        rows_tokens: list[list[object]] = []
        in_notes = False
        for token in self.tokens:
            kind = token.type
            if kind in ("footnote_block_open", "footnote_block_close"):
                in_notes = kind == "footnote_block_open"  # Written as Word's footnotes, at their references.
            elif in_notes:
                continue
            elif kind == "heading_open":
                heading = int(token.tag[1])
            elif kind == "heading_close":
                heading = 0
            elif kind in ("bullet_list_open", "ordered_list_open"):
                lists.append("Number" if kind == "ordered_list_open" else "Bullet")
            elif kind in ("bullet_list_close", "ordered_list_close"):
                lists.pop()
            elif kind == "blockquote_open":
                quote += 1
            elif kind == "blockquote_close":
                quote -= 1
            elif kind == "fence" and is_chart(token.info):
                self._chart(token.content)
            elif kind == "fence" and is_flowchart(token.info):
                self._flowchart(token.content)
            elif kind in ("fence", "code_block"):
                for line in token.content.rstrip("\n").split("\n"):
                    paragraph = self.document.add_paragraph(style="No Spacing")
                    run = paragraph.add_run(line)
                    run.font.name = _MONOSPACE
            elif kind == "hr":
                self._rule()
            elif kind == "table_open":
                table, rows_tokens = [], []
            elif kind == "tr_open" and table is not None:
                rows_tokens.append([])
            elif kind in ("th_open", "td_open"):
                cell_token = [kind]
            elif kind == "table_close" and table is not None:
                self._table(rows_tokens)
                table = None
            elif kind == "inline":
                if table is not None and cell_token:
                    rows_tokens[-1].append((cell_token[0], token))
                    cell_token = []
                    continue
                text = "".join(child.content for child in token.children or () if child.type == "text").strip()
                only = _PLACEHOLDER.fullmatch(text)
                if only is not None and self.zones[int(only.group(1))].display and len(token.children or ()) == 1:
                    self._display(int(only.group(1)))
                    continue
                if heading:
                    paragraph = self.document.add_heading(level=min(heading, 9))
                elif lists:
                    level = len(lists)
                    style = f"List {lists[-1]}" + (f" {level}" if 1 < level <= 3 else "")
                    paragraph = self.document.add_paragraph(style=style)
                elif quote:
                    paragraph = self.document.add_paragraph(style="Quote")
                else:
                    paragraph = self.document.add_paragraph()
                self.inline(paragraph, token)

    def _table(self, rows: list[list[tuple[str, object]]]) -> None:
        columns = max((len(row) for row in rows), default=1)
        table = self.document.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column, (kind, token) in enumerate(row):
                paragraph = table.cell(row_index, column).paragraphs[0]
                self.inline(paragraph, token)
                if kind == "th_open":
                    for run in paragraph.runs:
                        run.bold = True
        self.document.add_paragraph()

    def _rule(self) -> None:
        paragraph = self.document.add_paragraph()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for name, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "auto")):
            bottom.set(qn(name), value)
        borders.append(bottom)
        paragraph._p.get_or_add_pPr().append(borders)


def _field(paragraph: object, instruction: str, placeholder: str = "") -> None:
    """A Word field (PAGE, TOC), its result filled in when Word updates the fields."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    run._r.append(code)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    if placeholder:
        run._r.append(_text_element(placeholder))
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _text_element(text: str) -> object:
    element = OxmlElement("w:t")
    element.text = text
    return element


def _set_up(document: Document, options: DocxOptions, title: str) -> None:
    properties = document.core_properties
    properties.title = options.title or title
    properties.author = options.author
    properties.subject = options.subject
    properties.keywords = ", ".join(options.keywords)
    properties.language = _language_tag(options.language)

    normal = document.styles["Normal"]
    normal.font.name = options.font
    normal.font.size = Pt(options.font_size)
    normal.paragraph_format.line_spacing = options.line_spacing
    if options.justify:
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # The document's language, for proofing and for screen readers.
    defaults = document.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        run_defaults = defaults.find(qn("w:rPrDefault"))
        run_properties = run_defaults.find(qn("w:rPr")) if run_defaults is not None else None
        if run_properties is not None:
            language = run_properties.find(qn("w:lang"))
            if language is None:
                language = OxmlElement("w:lang")
                run_properties.append(language)
            language.set(qn("w:val"), _language_tag(options.language))

    section = document.sections[0]
    width, height = PAPERS.get(options.paper, PAPERS["a4"])
    if options.landscape:
        width, height = height, width
        section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(width), Mm(height)
    top, right, bottom, left = options.margins_mm
    section.top_margin, section.right_margin = Mm(top), Mm(right)
    section.bottom_margin, section.left_margin = Mm(bottom), Mm(left)


def _running_heads(section: object, options: DocxOptions, running: HeaderFooter) -> None:
    """The section's header and footer: the document's text at the left, the page number centred."""
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = running.header
    width = section.page_width - section.left_margin - section.right_margin
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(width // 2, WD_TAB_ALIGNMENT.CENTER)
    if running.footer:
        footer.add_run(running.footer)
    if options.page_numbers:
        if running.footer:
            footer.add_run("\t")
        else:
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(footer, "PAGE", "1")


def docx_document(pages: Sequence[ExportPage], title: str, omml_of: OmmlOf, options: DocxOptions) -> bytes:
    """The whole document as DOCX bytes: a notebook's pages under their titles, each on a new page.

    A page whose header or footer differs from the one before starts a section of its own.
    """
    document = Document()
    _set_up(document, options, title)
    heads = [header_and_footer(page.source) for page in pages]
    _running_heads(document.sections[0], options, heads[0] if heads else HeaderFooter())
    footnotes = _Footnotes(document)
    if options.table_of_contents:
        document.add_heading("Contents", level=1)
        _field(document.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "The table of contents is filled in by Word.")
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")  # Word offers to fill it in on opening.
        document.settings.element.append(update)
        document.add_page_break()
    for number, page in enumerate(pages):
        if len(pages) > 1:
            if number and heads[number] != heads[number - 1]:
                _running_heads(document.add_section(WD_SECTION.NEW_PAGE), options, heads[number])
            elif number:
                document.add_page_break()
            document.add_heading(page.title, level=1)
        _Writer(document, page.source, for_page(omml_of, page), footnotes, page.folder).body()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
