import io
import unittest
import zipfile

from docx import Document

from mathessible.documents.docx_export import DocxOptions, docx_document
from mathessible.documents.html_export import ExportPage
from mathessible.documents.omml_math import OmmlEquation

DOCUMENT = """# Roots

The roots of $x$ are *given* by **this**, see [here](https://example.com) and `code`:

$$y$$

$$z$$

- one
  - nested
1. first

| a | b |
|---|---|
| 1 | $x$ |

> quoted

```
a = 1
```
"""


def omml_of(source: str) -> OmmlEquation | None:
    if source == "x":
        return OmmlEquation('<m:r><m:t xml:space="preserve">x</m:t></m:r>')
    if source == "z":
        return OmmlEquation('<m:r><m:t xml:space="preserve">z</m:t></m:r>', "(1)")
    return OmmlEquation('<m:r><m:t xml:space="preserve">y</m:t></m:r>')


class DocxExportTests(unittest.TestCase):
    def export(self, **options: object) -> tuple[Document, str]:
        data = docx_document([ExportPage("Roots", DOCUMENT)], "Roots", omml_of, DocxOptions(**options))
        xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
        return Document(io.BytesIO(data)), xml

    def test_equations_are_words_own_inline_and_display(self) -> None:
        _document, xml = self.export()
        self.assertEqual(xml.count("<m:oMath>") + xml.count("<m:oMath "), 4)  # x twice (text, table), y, z.
        self.assertEqual(xml.count("<m:oMathPara>"), 2)
        self.assertIn("<m:t>#</m:t>", xml)  # z's number, as Word numbers an equation.

    def test_structure_uses_words_styles(self) -> None:
        document, xml = self.export()
        styles = [paragraph.style.name for paragraph in document.paragraphs]
        for expected in ("Heading 1", "List Bullet", "List Bullet 2", "List Number", "Quote", "No Spacing"):
            self.assertIn(expected, styles)
        self.assertEqual(len(document.tables), 1)
        self.assertIn("w:hyperlink", xml)
        self.assertTrue(any(run.bold for paragraph in document.paragraphs for run in paragraph.runs))

    def test_properties_page_and_language(self) -> None:
        document, xml = self.export(
            title="T",
            author="Ana",
            subject="S",
            keywords=("k1", "k2"),
            language="ro",
            paper="us-letter",
            landscape=True,
            margins_mm=(10, 20, 30, 40),
            page_numbers=True,
            table_of_contents=True,
        )
        properties = document.core_properties
        self.assertEqual((properties.title, properties.author, properties.subject), ("T", "Ana", "S"))
        self.assertEqual((properties.keywords, properties.language), ("k1, k2", "ro-RO"))
        section = document.sections[0]
        self.assertGreater(section.page_width, section.page_height)
        self.assertEqual(round(section.top_margin.mm), 10)
        self.assertIn("PAGE", "".join(section.footer.paragraphs[0]._p.itertext()))
        self.assertIn("TOC", xml)
        self.assertIn('w:val="ro-RO"', document.styles.element.xml)

    def test_footnotes_are_words_own_with_their_equations(self) -> None:
        markdown = "Roots[^a] of $x$, again[^a], and[^b].\n\n[^a]: Where $x$ is [real](https://e.com).\n\n[^b]: Two.\n"
        data = docx_document([ExportPage("Notes", markdown)], "Notes", omml_of, DocxOptions())
        archive = zipfile.ZipFile(io.BytesIO(data))
        body = archive.read("word/document.xml").decode("utf-8")
        notes = archive.read("word/footnotes.xml").decode("utf-8")
        self.assertEqual(body.count("<w:footnoteReference"), 2)
        self.assertIn('<w:footnoteReference w:id="1"/>', body)
        self.assertIn(">1</w:t>", body)  # The note referenced again: its number.
        self.assertIn('w:type="separator"', notes)
        self.assertEqual(notes.count("<w:footnoteRef/>"), 2)
        self.assertIn("<m:oMath>", notes)
        self.assertIn("https://e.com", archive.read("word/_rels/footnotes.xml.rels").decode("utf-8"))
        self.assertNotIn("[^", body)
        styles = Document(io.BytesIO(data)).styles
        self.assertTrue(styles["footnote reference"].font.superscript)

    def test_the_documents_header_and_footer_and_a_page_with_none(self) -> None:
        pages = [ExportPage("A", "---\nheader: Calculus\nfooter: Chapter 2\n---\nText"), ExportPage("B", "Text")]
        data = docx_document(pages, "Notebook", omml_of, DocxOptions(page_numbers=True))
        document = Document(io.BytesIO(data))
        first, second = document.sections
        self.assertEqual(first.header.paragraphs[0].text, "Calculus")
        self.assertEqual(first.footer.paragraphs[0].text, "Chapter 2\t1")  # The page number after a tab.
        self.assertEqual(second.header.paragraphs[0].text, "")
        self.assertEqual(second.footer.paragraphs[0].text, "1")
        self.assertNotIn("header:", "".join(paragraph.text for paragraph in document.paragraphs))


if __name__ == "__main__":
    unittest.main()
